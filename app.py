from flask import Flask, render_template, request, redirect, flash, session, make_response, g, has_request_context
from werkzeug.security import generate_password_hash, check_password_hash
from functools import wraps
import psycopg2
import psycopg2.extras
import datetime
import json
import csv
import io
import calendar
import os

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "expense-manager-secret-key")

CURRENCY_SYMBOLS = {
    "INR": "₹",
    "USD": "$",
    "EUR": "€",
    "GBP": "£",
    "JPY": "¥"
}

DEFAULT_CATEGORY_ICONS = {
    "Food": "🍔",
    "Shopping": "🛍️",
    "Travel": "🚗",
    "Bills": "💡",
    "Entertainment": "🎮",
    "Health": "❤️",
    "Education": "📚",
    "Other": "📦"
}

DEFAULT_PAYMENT_METHODS = (
    ("Cash", "💵"),
    ("UPI", "📱"),
    ("Bank Transfer", "🏦"),
    ("Card", "💳"),
    ("Other", "💰")
)

_db_initialized = False


def _create_raw_connection():
    url = os.environ.get("DATABASE_URL")
    if not url:
        raise RuntimeError("DATABASE_URL environment variable is not configured. Please set DATABASE_URL in your environment or Vercel project settings.")

    if url.startswith("postgres://"):
        url = url.replace("postgres://", "postgresql://", 1)

    return psycopg2.connect(url, cursor_factory=psycopg2.extras.DictCursor)


class RequestConnection:
    """Thin wrapper around a psycopg2 connection that survives route-level
    close() calls by reconnecting on next use. This lets every helper inside a
    single HTTP request reuse one Neon connection instead of opening several
    (each costs a TLS handshake + serverless round trip)."""

    def __init__(self, raw):
        self._raw = raw

    def _ensure(self):
        if self._raw.closed:
            self._raw = _create_raw_connection()
        return self._raw

    @property
    def closed(self):
        return self._raw.closed

    def cursor(self, *args, **kwargs):
        return self._ensure().cursor(*args, **kwargs)

    def commit(self):
        return self._ensure().commit()

    def rollback(self):
        return self._ensure().rollback()

    def close(self):
        # Keep the request-scoped connection alive so later helpers reuse it.
        # The real close happens in teardown_appcontext.
        pass

    def force_close(self):
        try:
            self._raw.close()
        except Exception:
            pass


@app.teardown_appcontext
def close_request_connection(exception):
    wrapper = g.pop("_db_conn", None)
    if wrapper is not None:
        wrapper.force_close()


def get_db_connection():
    """Return a single connection shared by everything in the current request."""
    global _db_initialized

    if not has_request_context():
        conn = _create_raw_connection()
        if not _db_initialized:
            try:
                init_db_with_conn(conn)
                _db_initialized = True
            except Exception as e:
                print(f"Error initializing DB schema: {e}")
        return conn

    wrapper = getattr(g, "_db_conn", None)
    if wrapper is None:
        wrapper = RequestConnection(_create_raw_connection())
        g._db_conn = wrapper

    if not _db_initialized:
        try:
            init_db_with_conn(wrapper._ensure())
            _db_initialized = True
        except Exception as e:
            print(f"Error initializing DB schema: {e}")

    return wrapper


def init_db_with_conn(conn):
    with conn.cursor() as cursor:
        # Create users table
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS users (
                id SERIAL PRIMARY KEY,
                username VARCHAR(100) NOT NULL UNIQUE,
                email VARCHAR(255) NOT NULL UNIQUE,
                password_hash TEXT NOT NULL,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)

        # Create expenses table
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS expenses (
                id SERIAL PRIMARY KEY,
                user_id INTEGER REFERENCES users(id) ON DELETE CASCADE,
                amount DOUBLE PRECISION NOT NULL,
                category VARCHAR(255) NOT NULL,
                description TEXT,
                date VARCHAR(50) NOT NULL,
                payment_method VARCHAR(100) DEFAULT 'Cash'
            )
        """)

        # Create income table
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS income (
                id SERIAL PRIMARY KEY,
                user_id INTEGER REFERENCES users(id) ON DELETE CASCADE,
                amount DOUBLE PRECISION NOT NULL,
                category VARCHAR(255) NOT NULL,
                description TEXT,
                date VARCHAR(50) NOT NULL,
                payment_method VARCHAR(100) DEFAULT 'Cash'
            )
        """)

        # Create user_settings table
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS user_settings (
                user_id INTEGER REFERENCES users(id) ON DELETE CASCADE,
                key VARCHAR(255) NOT NULL,
                value TEXT NOT NULL,
                PRIMARY KEY (user_id, key)
            )
        """)

        # Create categories table
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS categories (
                id SERIAL PRIMARY KEY,
                user_id INTEGER REFERENCES users(id) ON DELETE CASCADE,
                name VARCHAR(255) NOT NULL,
                icon VARCHAR(50) DEFAULT '📦'
            )
        """)

        # Create payment_methods table
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS payment_methods (
                id SERIAL PRIMARY KEY,
                user_id INTEGER REFERENCES users(id) ON DELETE CASCADE,
                name VARCHAR(255) NOT NULL,
                icon VARCHAR(50) DEFAULT '💳'
            )
        """)

        # Alter table migrations if columns exist from previous runs
        cursor.execute("ALTER TABLE expenses ADD COLUMN IF NOT EXISTS user_id INTEGER REFERENCES users(id) ON DELETE CASCADE")
        cursor.execute("ALTER TABLE income ADD COLUMN IF NOT EXISTS user_id INTEGER REFERENCES users(id) ON DELETE CASCADE")
        cursor.execute("ALTER TABLE categories ADD COLUMN IF NOT EXISTS user_id INTEGER REFERENCES users(id) ON DELETE CASCADE")
        cursor.execute("ALTER TABLE payment_methods ADD COLUMN IF NOT EXISTS user_id INTEGER REFERENCES users(id) ON DELETE CASCADE")

        # Drop old single-column unique constraint on name if it exists from previous migrations
        try:
            cursor.execute("ALTER TABLE categories DROP CONSTRAINT IF EXISTS categories_name_key")
        except Exception:
            pass
            
        try:
            cursor.execute("ALTER TABLE payment_methods DROP CONSTRAINT IF EXISTS payment_methods_name_key")
        except Exception:
            pass

        conn.commit()


def ensure_user_defaults(conn, user_id):
    with conn.cursor() as cursor:
        # Default user settings
        cursor.execute("INSERT INTO user_settings (user_id, key, value) VALUES (%s, 'currency', 'INR') ON CONFLICT (user_id, key) DO NOTHING", (user_id,))
        cursor.execute("INSERT INTO user_settings (user_id, key, value) VALUES (%s, 'monthly_budget', '10000.00') ON CONFLICT (user_id, key) DO NOTHING", (user_id,))

        # Categories: seed defaults once; otherwise realign existing default icons
        # with a single idempotent UPDATE instead of one query per category.
        cursor.execute("SELECT COUNT(*) FROM categories WHERE user_id = %s", (user_id,))
        if cursor.fetchone()[0] == 0:
            cursor.executemany(
                "INSERT INTO categories (user_id, name, icon) VALUES (%s, %s, %s)",
                [(user_id, name, icon) for name, icon in DEFAULT_CATEGORY_ICONS.items()]
            )
        else:
            placeholders = ", ".join(["(%s, %s)"] * len(DEFAULT_CATEGORY_ICONS))
            params = [icon for pair in DEFAULT_CATEGORY_ICONS.items() for icon in pair]
            params.append(user_id)
            cursor.execute(
                f"""
                UPDATE categories c
                SET icon = v.icon
                FROM (VALUES {placeholders}) AS v(name, icon)
                WHERE c.user_id = %s AND c.name = v.name AND c.icon IS DISTINCT FROM v.icon
                """,
                params
            )

        # Payment methods: seed defaults once.
        cursor.execute("SELECT COUNT(*) FROM payment_methods WHERE user_id = %s", (user_id,))
        if cursor.fetchone()[0] == 0:
            cursor.executemany(
                "INSERT INTO payment_methods (user_id, name, icon) VALUES (%s, %s, %s)",
                [(user_id, name, icon) for name, icon in DEFAULT_PAYMENT_METHODS]
            )

        conn.commit()


def create_default_user_data(conn, user_id):
    ensure_user_defaults(conn, user_id)


def init_db():
    if os.environ.get("DATABASE_URL"):
        try:
            conn = get_db_connection()
            conn.close()
        except Exception as e:
            print(f"init_db error: {e}")


def get_setting(user_id, key, default):
    if not user_id:
        return default
    connection = get_db_connection()
    cursor = connection.cursor()
    cursor.execute("SELECT value FROM user_settings WHERE user_id = %s AND key = %s", (user_id, key))
    row = cursor.fetchone()
    connection.close()
    if row:
        return row[0]
    return default


def set_setting(user_id, key, value):
    if not user_id:
        return
    connection = get_db_connection()
    cursor = connection.cursor()
    cursor.execute(
        "INSERT INTO user_settings (user_id, key, value) VALUES (%s, %s, %s) ON CONFLICT (user_id, key) DO UPDATE SET value = EXCLUDED.value",
        (user_id, key, str(value))
    )
    connection.commit()
    connection.close()


def get_currency_info(user_id):
    currency_code = get_setting(user_id, "currency", "INR")
    currency_symbol = CURRENCY_SYMBOLS.get(currency_code, "₹")
    return currency_code, currency_symbol


def get_currency_symbol_and_budget(user_id):
    """Fetch currency symbol + monthly budget in ONE round trip. These two
    settings are shown together in nearly every page's sidebar, so combining
    them halves the per-page settings queries."""
    currency_symbol = "₹"
    budget = "10000.00"
    if user_id:
        connection = get_db_connection()
        cursor = connection.cursor()
        cursor.execute(
            "SELECT key, value FROM user_settings WHERE user_id = %s AND key IN ('currency', 'monthly_budget')",
            (user_id,)
        )
        for row in cursor.fetchall():
            if row["key"] == "currency":
                currency_symbol = CURRENCY_SYMBOLS.get(row["value"], "₹")
            elif row["key"] == "monthly_budget":
                budget = row["value"]
    return currency_symbol, float(budget)


def get_sorted_payment_methods(cursor, user_id):
    cursor.execute("SELECT * FROM payment_methods WHERE user_id = %s", (user_id,))
    rows = cursor.fetchall()
    order_map = {"Cash": 1, "UPI": 2, "Bank Transfer": 3, "Card": 4, "Other": 5}
    return sorted(rows, key=lambda x: order_map.get(x["name"], 99))


def login_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if "user_id" not in session:
            flash("Please log in to access Expense Manager.", "error")
            return redirect("/login")

        user_id = session["user_id"]
        # Seed per-user defaults at most once per session. This avoids running
        # several database queries on every single request just to re-verify
        # defaults that only need to be created for new users.
        if not session.get("_defaults_seeded"):
            try:
                conn = get_db_connection()
                ensure_user_defaults(conn, user_id)
                session["_defaults_seeded"] = True
            except Exception as e:
                print(f"ensure_user_defaults error: {e}")

        return f(*args, **kwargs)
    return decorated_function


def get_filter_dates():
    now = datetime.datetime.now()
    today = now.date()
    
    filter_range = request.args.get("range", "").strip().lower()
    if filter_range:
        session["filter_range"] = filter_range
    else:
        filter_range = session.get("filter_range", "this_month")
        
    if filter_range not in ["today", "this_week", "this_month", "last_month", "this_year", "custom"]:
        filter_range = "this_month"
        
    selected_month = request.args.get("month", "").strip()
    if selected_month:
        session["selected_month"] = selected_month
    else:
        selected_month = session.get("selected_month", now.strftime("%Y-%m"))
        
    try:
        parsed_date = datetime.datetime.strptime(selected_month, "%Y-%m")
    except ValueError:
        selected_month = now.strftime("%Y-%m")
        parsed_date = now

    if filter_range == "today":
        start_date = today.strftime("%Y-%m-%d")
        end_date = today.strftime("%Y-%m-%d")
        range_label = "Today"
    elif filter_range == "this_week":
        start_of_week = today - datetime.timedelta(days=today.weekday())
        end_of_week = start_of_week + datetime.timedelta(days=6)
        start_date = start_of_week.strftime("%Y-%m-%d")
        end_date = end_of_week.strftime("%Y-%m-%d")
        range_label = "This Week"
    elif filter_range == "this_month":
        start_date = today.replace(day=1).strftime("%Y-%m-%d")
        next_month = today.replace(day=28) + datetime.timedelta(days=4)
        last_day = next_month - datetime.timedelta(days=next_month.day)
        end_date = last_day.strftime("%Y-%m-%d")
        range_label = today.strftime("%B %Y")
    elif filter_range == "last_month":
        first_day_this_month = today.replace(day=1)
        last_day_last_month = first_day_this_month - datetime.timedelta(days=1)
        first_day_last_month = last_day_last_month.replace(day=1)
        start_date = first_day_last_month.strftime("%Y-%m-%d")
        end_date = last_day_last_month.strftime("%Y-%m-%d")
        range_label = last_day_last_month.strftime("%B %Y")
    elif filter_range == "this_year":
        start_date = today.replace(month=1, day=1).strftime("%Y-%m-%d")
        end_date = today.replace(month=12, day=31).strftime("%Y-%m-%d")
        range_label = today.strftime("%Y")
    else: # custom
        first_day = parsed_date.date().replace(day=1)
        next_month = first_day.replace(day=28) + datetime.timedelta(days=4)
        last_day = next_month - datetime.timedelta(days=next_month.day)
        start_date = first_day.strftime("%Y-%m-%d")
        end_date = last_day.strftime("%Y-%m-%d")
        range_label = parsed_date.strftime("%B %Y")
        
    return start_date, end_date, filter_range, selected_month, range_label


# =========================================
# AUTH ROUTES
# =========================================

@app.route("/login", methods=["GET", "POST"])
def login():
    if "user_id" in session:
        return redirect("/")

    if request.method == "POST":
        username_or_email = request.form.get("username_or_email", "").strip()
        password = request.form.get("password", "")

        if not username_or_email or not password:
            flash("Please enter both username/email and password.", "error")
            return render_template("login.html")

        connection = get_db_connection()
        cursor = connection.cursor()
        cursor.execute(
            "SELECT * FROM users WHERE LOWER(username) = LOWER(%s) OR LOWER(email) = LOWER(%s)",
            (username_or_email, username_or_email)
        )
        user = cursor.fetchone()
        connection.close()

        if user and check_password_hash(user["password_hash"], password):
            session["user_id"] = user["id"]
            session["username"] = user["username"]
            session["email"] = user["email"]
            flash(f"Welcome back, {user['username']}! 👋", "success")
            return redirect("/")
        else:
            flash("Invalid username/email or password.", "error")

    return render_template("login.html")


@app.route("/signup", methods=["GET", "POST"])
def signup():
    if "user_id" in session:
        return redirect("/")

    if request.method == "POST":
        username = request.form.get("username", "").strip()
        email = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")
        confirm_password = request.form.get("confirm_password", "")

        if not username or not email or not password:
            flash("Please fill in all required fields.", "error")
            return render_template("signup.html")

        if len(username) < 3:
            flash("Username must be at least 3 characters long.", "error")
            return render_template("signup.html")

        if "@" not in email:
            flash("Please enter a valid email address.", "error")
            return render_template("signup.html")

        if len(password) < 6:
            flash("Password must be at least 6 characters long.", "error")
            return render_template("signup.html")

        if password != confirm_password:
            flash("Passwords do not match.", "error")
            return render_template("signup.html")

        connection = get_db_connection()
        cursor = connection.cursor()

        # Check existing user
        cursor.execute(
            "SELECT COUNT(*) FROM users WHERE LOWER(username) = LOWER(%s) OR LOWER(email) = LOWER(%s)",
            (username, email)
        )
        if cursor.fetchone()[0] > 0:
            connection.close()
            flash("Username or Email is already registered. Please log in.", "error")
            return render_template("signup.html")

        # Create user
        pwd_hash = generate_password_hash(password)
        cursor.execute(
            "INSERT INTO users (username, email, password_hash) VALUES (%s, %s, %s) RETURNING id",
            (username, email, pwd_hash)
        )
        user_id = cursor.fetchone()[0]
        connection.commit()

        # Seed initial user settings, categories, payment methods
        create_default_user_data(connection, user_id)
        connection.close()

        # Auto log in user
        session["user_id"] = user_id
        session["username"] = username
        session["email"] = email
        session["_defaults_seeded"] = True
        flash(f"Account created successfully! Welcome, {username}! 🎉", "success")
        return redirect("/")

    return render_template("signup.html")


@app.route("/logout")
def logout():
    session.clear()
    flash("You have been logged out.", "info")
    return redirect("/login")


# =========================================
# APP MAIN DASHBOARD & VIEW ROUTES
# =========================================

@app.route("/")
@login_required
def home():
    user_id = session["user_id"]
    connection = get_db_connection()
    cursor = connection.cursor()

    # Calculate Total Expenses and Total Income (All-time for this user)
    cursor.execute("""
        SELECT
            (SELECT COALESCE(SUM(amount), 0) FROM expenses WHERE user_id = %s),
            (SELECT COALESCE(SUM(amount), 0) FROM income WHERE user_id = %s)
    """, (user_id, user_id))
    sums_row = cursor.fetchone()
    total_expenses = float(sums_row[0] or 0)
    total_income = float(sums_row[1] or 0)

    # Calculate Total Balance
    balance = total_income - total_expenses

    # Get Date Filter Range bounds using helper
    start_date, end_date, filter_range, selected_month, range_label = get_filter_dates()

    # Calculate Monthly Expenses + Spent Today in a single query (one round trip)
    now = datetime.datetime.now()
    today_str = now.strftime("%Y-%m-%d")
    cursor.execute("""
        SELECT
            COALESCE(SUM(amount), 0) AS month_total,
            COALESCE(SUM(amount) FILTER (WHERE date = %s), 0) AS today_total
        FROM expenses
        WHERE user_id = %s AND date BETWEEN %s AND %s
    """, (today_str, user_id, start_date, end_date))
    range_row = cursor.fetchone()
    monthly_expenses = range_row["month_total"]
    today_expenses = range_row["today_total"]

    # Get Daily Expenses for Chart.js
    cursor.execute("""
        SELECT date, SUM(amount) AS total 
        FROM expenses 
        WHERE user_id = %s AND date BETWEEN %s AND %s
        GROUP BY date 
        ORDER BY date ASC
    """, (user_id, start_date, end_date))
    chart_rows = cursor.fetchall()

    # Format Chart.js Data
    chart_labels = []
    chart_values = []
    months_short = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

    for row in chart_rows:
        try:
            d = datetime.datetime.strptime(row["date"], "%Y-%m-%d")
            day_label = f"{d.day} {months_short[d.month - 1]}"
        except ValueError:
            day_label = row["date"]
        chart_labels.append(day_label)
        chart_values.append(float(row["total"]))

    # Fallbacks if chart data is empty
    if not chart_labels:
        try:
            m_idx = int(selected_month.split('-')[1])
            m_short = months_short[m_idx - 1]
        except (ValueError, IndexError):
            m_short = "Aug"
        chart_labels = [f"1 {m_short}", f"5 {m_short}", f"10 {m_short}", f"15 {m_short}", f"20 {m_short}", f"25 {m_short}", f"31 {m_short}"]
        chart_values = [0, 0, 0, 0, 0, 0, 0]

    chart_labels_json = json.dumps(chart_labels)
    chart_values_json = json.dumps(chart_values)

    # Fetch dynamic currency + budget settings (single round trip)
    currency_symbol, budget_limit = get_currency_symbol_and_budget(user_id)
    budget_percent = min(100, int((monthly_expenses / budget_limit) * 100)) if budget_limit > 0 else 0

    # Get Expenses for the SELECTED date range
    cursor.execute("""
        SELECT e.*, c.icon AS category_icon
        FROM expenses e
        LEFT JOIN categories c ON (e.category = c.name AND c.user_id = e.user_id)
        WHERE e.user_id = %s AND e.date BETWEEN %s AND %s
        ORDER BY e.date DESC, e.id DESC
    """, (user_id, start_date, end_date))
    expenses = cursor.fetchall()

    connection.close()

    return render_template(
        "dashboard.html",
        total_expenses=total_expenses,
        total_income=total_income,
        balance=balance,
        monthly_expenses=monthly_expenses,
        today_expenses=today_expenses,
        selected_month=selected_month,
        selected_month_name=range_label,
        filter_range=filter_range,
        chart_labels_json=chart_labels_json,
        chart_values_json=chart_values_json,
        budget_limit=budget_limit,
        budget_percent=budget_percent,
        currency_symbol=currency_symbol,
        expenses=expenses
    )


@app.route("/add-expense", methods=["GET", "POST"])
@login_required
def add_expense():
    user_id = session["user_id"]
    if request.method == "POST":
        amount = request.form["amount"]
        category = request.form["category"]
        description = request.form["description"]
        date = request.form["date"]
        payment_method = request.form.get("payment_method", "Cash")

        connection = get_db_connection()
        cursor = connection.cursor()
        cursor.execute("""
            INSERT INTO expenses
            (user_id, amount, category, description, date, payment_method)
            VALUES (%s, %s, %s, %s, %s, %s)
        """, (user_id, amount, category, description, date, payment_method))
        connection.commit()
        connection.close()

        flash("Expense added successfully! 💸", "success")
        return redirect("/")

    # Fetch options dynamically
    connection = get_db_connection()
    cursor = connection.cursor()
    
    cursor.execute("SELECT * FROM categories WHERE user_id = %s ORDER BY name ASC", (user_id,))
    categories = cursor.fetchall()
    
    payment_methods = get_sorted_payment_methods(cursor, user_id)

    # Monthly overview sidebar metrics
    now = datetime.datetime.now()
    current_month_str = now.strftime("%Y-%m")
    current_month_name = now.strftime("%B %Y")
    
    cursor.execute("SELECT SUM(amount) AS total FROM expenses WHERE user_id = %s AND substr(date, 1, 7) = %s", (user_id, current_month_str))
    monthly_expenses = cursor.fetchone()[0] or 0
    connection.close()

    currency_symbol, budget_limit = get_currency_symbol_and_budget(user_id)
    budget_percent = min(100, int((monthly_expenses / budget_limit) * 100)) if budget_limit > 0 else 0

    return render_template(
        "add_expense.html",
        current_month_name=current_month_name,
        monthly_expenses=monthly_expenses,
        budget_limit=budget_limit,
        budget_percent=budget_percent,
        categories=categories,
        payment_methods=payment_methods,
        currency_symbol=currency_symbol
    )


@app.route("/add-income", methods=["GET", "POST"])
@login_required
def add_income():
    user_id = session["user_id"]
    if request.method == "POST":
        amount = request.form["amount"]
        category = request.form["category"]
        description = request.form["description"]
        date = request.form["date"]
        payment_method = request.form.get("payment_method", "Cash")

        connection = get_db_connection()
        cursor = connection.cursor()
        cursor.execute("""
            INSERT INTO income
            (user_id, amount, category, description, date, payment_method)
            VALUES (%s, %s, %s, %s, %s, %s)
        """, (user_id, amount, category, description, date, payment_method))
        connection.commit()
        connection.close()

        flash("Income added successfully! 💰", "success")
        return redirect("/")

    # Fetch options dynamically
    connection = get_db_connection()
    cursor = connection.cursor()
    
    cursor.execute("SELECT * FROM categories WHERE user_id = %s ORDER BY name ASC", (user_id,))
    categories = cursor.fetchall()
    
    payment_methods = get_sorted_payment_methods(cursor, user_id)

    # Monthly overview sidebar metrics
    now = datetime.datetime.now()
    current_month_str = now.strftime("%Y-%m")
    current_month_name = now.strftime("%B %Y")
    
    cursor.execute("SELECT SUM(amount) AS total FROM expenses WHERE user_id = %s AND substr(date, 1, 7) = %s", (user_id, current_month_str))
    monthly_expenses = cursor.fetchone()[0] or 0
    connection.close()

    currency_symbol, budget_limit = get_currency_symbol_and_budget(user_id)
    budget_percent = min(100, int((monthly_expenses / budget_limit) * 100)) if budget_limit > 0 else 0

    return render_template(
        "add_income.html",
        current_month_name=current_month_name,
        monthly_expenses=monthly_expenses,
        budget_limit=budget_limit,
        budget_percent=budget_percent,
        categories=categories,
        payment_methods=payment_methods,
        currency_symbol=currency_symbol
    )


@app.route("/delete-expense/<int:expense_id>", methods=["POST"])
@login_required
def delete_expense(expense_id):
    user_id = session["user_id"]
    connection = get_db_connection()
    cursor = connection.cursor()
    cursor.execute("DELETE FROM expenses WHERE id = %s AND user_id = %s", (expense_id, user_id))
    connection.commit()
    connection.close()
    flash("Expense deleted.", "info")
    return redirect(request.referrer or "/")


@app.route("/delete-income/<int:income_id>", methods=["POST"])
@login_required
def delete_income(income_id):
    user_id = session["user_id"]
    connection = get_db_connection()
    cursor = connection.cursor()
    cursor.execute("DELETE FROM income WHERE id = %s AND user_id = %s", (income_id, user_id))
    connection.commit()
    connection.close()
    flash("Income deleted.", "info")
    return redirect(request.referrer or "/")


# =========================================
# SETTINGS PAGE ROUTES
# =========================================

@app.route("/settings")
@login_required
def settings():
    user_id = session["user_id"]
    connection = get_db_connection()
    cursor = connection.cursor()

    # Fetch settings data
    currency_code, currency_symbol = get_currency_info(user_id)
    budget_limit = float(get_setting(user_id, "monthly_budget", "10000.00"))

    # Fetch monthly spending aggregates
    now = datetime.datetime.now()
    current_month_str = now.strftime("%Y-%m")
    current_month_name = now.strftime("%B %Y")
    
    cursor.execute("SELECT SUM(amount) AS total FROM expenses WHERE user_id = %s AND substr(date, 1, 7) = %s", (user_id, current_month_str))
    monthly_expenses = cursor.fetchone()[0] or 0

    # Calculate remaining budget and percentage
    remaining_budget = budget_limit - monthly_expenses
    budget_percent = (monthly_expenses / budget_limit) * 100 if budget_limit > 0 else 0
    budget_percent_capped = min(100, int(budget_percent))

    # Fetch categories and payment methods
    cursor.execute("SELECT * FROM categories WHERE user_id = %s ORDER BY name ASC", (user_id,))
    categories = cursor.fetchall()

    payment_methods = get_sorted_payment_methods(cursor, user_id)

    connection.close()

    return render_template(
        "settings.html",
        currency_code=currency_code,
        currency_symbol=currency_symbol,
        budget_limit=budget_limit,
        monthly_expenses=monthly_expenses,
        remaining_budget=remaining_budget,
        budget_percent=budget_percent,
        budget_percent_capped=budget_percent_capped,
        current_month_name=current_month_name,
        categories=categories,
        payment_methods=payment_methods
    )


@app.route("/settings/currency", methods=["POST"])
@login_required
def update_currency():
    user_id = session["user_id"]
    new_currency = request.form.get("currency")
    if new_currency in CURRENCY_SYMBOLS:
        set_setting(user_id, "currency", new_currency)
        flash("Default currency updated successfully.", "success")
    else:
        flash("Invalid currency selection.", "error")
    return redirect("/settings")


@app.route("/settings/budget", methods=["POST"])
@login_required
def update_budget():
    user_id = session["user_id"]
    try:
        budget_val = float(request.form["monthly_budget"])
        if budget_val < 0:
            flash("Budget amount cannot be negative.", "error")
        else:
            set_setting(user_id, "monthly_budget", f"{budget_val:.2f}")
            flash("Monthly budget updated successfully.", "success")
    except ValueError:
        flash("Invalid budget format. Please enter a valid number.", "error")
    return redirect("/settings")


@app.route("/settings/category/add", methods=["POST"])
@login_required
def add_category():
    user_id = session["user_id"]
    name = request.form.get("name", "").strip()
    icon = request.form.get("icon", "📦").strip()

    if not name:
        flash("Category name cannot be empty.", "error")
        return redirect("/settings")

    try:
        connection = get_db_connection()
        cursor = connection.cursor()
        cursor.execute("INSERT INTO categories (user_id, name, icon) VALUES (%s, %s, %s)", (user_id, name, icon))
        connection.commit()
        connection.close()
        flash(f"Category '{name}' added successfully.", "success")
    except psycopg2.IntegrityError:
        flash("A category with this name already exists.", "error")
    
    return redirect("/settings")


@app.route("/settings/category/edit/<int:cat_id>", methods=["POST"])
@login_required
def edit_category(cat_id):
    user_id = session["user_id"]
    name = request.form.get("name", "").strip()
    icon = request.form.get("icon", "📦").strip()

    if not name:
        flash("Category name cannot be empty.", "error")
        return redirect("/settings")

    connection = get_db_connection()
    cursor = connection.cursor()
    
    # Get original category name
    cursor.execute("SELECT name FROM categories WHERE id = %s AND user_id = %s", (cat_id, user_id))
    orig_row = cursor.fetchone()
    
    if orig_row:
        orig_name = orig_row[0]
        try:
            cursor.execute("UPDATE categories SET name = %s, icon = %s WHERE id = %s AND user_id = %s", (name, icon, cat_id, user_id))
            
            if orig_name != name:
                cursor.execute("UPDATE expenses SET category = %s WHERE category = %s AND user_id = %s", (name, orig_name, user_id))
                cursor.execute("UPDATE income SET category = %s WHERE category = %s AND user_id = %s", (name, orig_name, user_id))
                
            connection.commit()
            flash("Category details updated successfully.", "success")
        except psycopg2.IntegrityError:
            flash("A category with this name already exists.", "error")
            
    connection.close()
    return redirect("/settings")


@app.route("/settings/category/delete/<int:cat_id>", methods=["POST"])
@login_required
def delete_category(cat_id):
    user_id = session["user_id"]
    connection = get_db_connection()
    cursor = connection.cursor()

    cursor.execute("SELECT name FROM categories WHERE id = %s AND user_id = %s", (cat_id, user_id))
    cat_row = cursor.fetchone()

    if cat_row:
        cat_name = cat_row[0]
        cursor.execute("SELECT COUNT(*) FROM expenses WHERE category = %s AND user_id = %s", (cat_name, user_id))
        expense_use = cursor.fetchone()[0]
        
        cursor.execute("SELECT COUNT(*) FROM income WHERE category = %s AND user_id = %s", (cat_name, user_id))
        income_use = cursor.fetchone()[0]

        if expense_use > 0 or income_use > 0:
            flash("This category is being used by existing transactions.", "error")
        else:
            cursor.execute("DELETE FROM categories WHERE id = %s AND user_id = %s", (cat_id, user_id))
            connection.commit()
            flash("Category deleted successfully.", "success")

    connection.close()
    return redirect("/settings")


@app.route("/settings/payment/add", methods=["POST"])
@login_required
def add_payment_method():
    user_id = session["user_id"]
    name = request.form.get("name", "").strip()
    icon = request.form.get("icon", "💳").strip()

    if not name:
        flash("Payment method name cannot be empty.", "error")
        return redirect("/settings")

    try:
        connection = get_db_connection()
        cursor = connection.cursor()
        cursor.execute("INSERT INTO payment_methods (user_id, name, icon) VALUES (%s, %s, %s)", (user_id, name, icon))
        connection.commit()
        connection.close()
        flash(f"Payment method '{name}' added successfully.", "success")
    except psycopg2.IntegrityError:
        flash("A payment method with this name already exists.", "error")

    return redirect("/settings")


@app.route("/settings/payment/edit/<int:pay_id>", methods=["POST"])
@login_required
def edit_payment_method(pay_id):
    user_id = session["user_id"]
    name = request.form.get("name", "").strip()
    icon = request.form.get("icon", "💳").strip()

    if not name:
        flash("Payment method name cannot be empty.", "error")
        return redirect("/settings")

    connection = get_db_connection()
    cursor = connection.cursor()

    cursor.execute("SELECT name FROM payment_methods WHERE id = %s AND user_id = %s", (pay_id, user_id))
    orig_row = cursor.fetchone()

    if orig_row:
        orig_name = orig_row[0]
        try:
            cursor.execute("UPDATE payment_methods SET name = %s, icon = %s WHERE id = %s AND user_id = %s", (name, icon, pay_id, user_id))
            
            if orig_name != name:
                cursor.execute("UPDATE expenses SET payment_method = %s WHERE payment_method = %s AND user_id = %s", (name, orig_name, user_id))
                cursor.execute("UPDATE income SET payment_method = %s WHERE payment_method = %s AND user_id = %s", (name, orig_name, user_id))
                
            connection.commit()
            flash("Payment method details updated successfully.", "success")
        except psycopg2.IntegrityError:
            flash("A payment method with this name already exists.", "error")

    connection.close()
    return redirect("/settings")


@app.route("/settings/payment/delete/<int:pay_id>", methods=["POST"])
@login_required
def delete_payment_method(pay_id):
    user_id = session["user_id"]
    connection = get_db_connection()
    cursor = connection.cursor()

    cursor.execute("SELECT name FROM payment_methods WHERE id = %s AND user_id = %s", (pay_id, user_id))
    pay_row = cursor.fetchone()

    if pay_row:
        pay_name = pay_row[0]
        cursor.execute("SELECT COUNT(*) FROM expenses WHERE payment_method = %s AND user_id = %s", (pay_name, user_id))
        expense_use = cursor.fetchone()[0]

        cursor.execute("SELECT COUNT(*) FROM income WHERE payment_method = %s AND user_id = %s", (pay_name, user_id))
        income_use = cursor.fetchone()[0]

        if expense_use > 0 or income_use > 0:
            flash("This payment method is being used by existing transactions.", "error")
        else:
            cursor.execute("DELETE FROM payment_methods WHERE id = %s AND user_id = %s", (pay_id, user_id))
            connection.commit()
            flash("Payment method deleted successfully.", "success")

    connection.close()
    return redirect("/settings")


# =========================================
# SIDEBAR NAVIGATION VIEW ROUTES
# =========================================

@app.route("/transactions")
@login_required
def transactions_view():
    user_id = session["user_id"]
    filter_type = request.args.get("type", "all").strip().lower()
    
    connection = get_db_connection()
    cursor = connection.cursor()
    
    # 1. Fetch filtered transactions
    if filter_type == "expense":
        query = """
            SELECT e.id, e.amount, e.category, e.description, e.date, e.payment_method, 'expense' AS type, c.icon AS category_icon
            FROM expenses e
            LEFT JOIN categories c ON (e.category = c.name AND c.user_id = e.user_id)
            WHERE e.user_id = %s
            ORDER BY e.date DESC, e.id DESC
        """
        params = (user_id,)
    elif filter_type == "income":
        query = """
            SELECT i.id, i.amount, i.category, i.description, i.date, i.payment_method, 'income' AS type, c.icon AS category_icon
            FROM income i
            LEFT JOIN categories c ON (i.category = c.name AND c.user_id = i.user_id)
            WHERE i.user_id = %s
            ORDER BY i.date DESC, i.id DESC
        """
        params = (user_id,)
    else:
        query = """
            SELECT t.id, t.amount, t.category, t.description, t.date, t.payment_method, t.type, c.icon AS category_icon
            FROM (
                SELECT id, user_id, amount, category, description, date, payment_method, 'expense' AS type FROM expenses WHERE user_id = %s
                UNION ALL
                SELECT id, user_id, amount, category, description, date, payment_method, 'income' AS type FROM income WHERE user_id = %s
            ) t
            LEFT JOIN categories c ON (t.category = c.name AND c.user_id = t.user_id)
            ORDER BY t.date DESC, t.id DESC
        """
        params = (user_id, user_id)
        
    cursor.execute(query, params)
    transactions = cursor.fetchall()
    
    # 2. Get dynamic overview metrics for sidebar
    now = datetime.datetime.now()
    current_month_str = now.strftime("%Y-%m")
    
    cursor.execute("SELECT SUM(amount) AS total FROM expenses WHERE user_id = %s AND substr(date, 1, 7) = %s", (user_id, current_month_str))
    monthly_expenses = cursor.fetchone()[0] or 0
    
    currency_symbol, budget_limit = get_currency_symbol_and_budget(user_id)
    budget_percent = min(100, int((monthly_expenses / budget_limit) * 100)) if budget_limit > 0 else 0
    
    connection.close()
    
    return render_template(
        "transactions.html",
        transactions=transactions,
        filter_type=filter_type,
        monthly_expenses=monthly_expenses,
        budget_limit=budget_limit,
        budget_percent=budget_percent,
        currency_symbol=currency_symbol
    )


@app.route("/analytics")
@login_required
def analytics():
    user_id = session["user_id"]
    connection = get_db_connection()
    cursor = connection.cursor()

    # Get Date Filter Range bounds using helper
    start_date, end_date, filter_range, selected_month, range_label = get_filter_dates()

    # 1. Total spent + total income in this active date range (one round trip)
    cursor.execute("""
        SELECT
            (SELECT COALESCE(SUM(amount), 0) FROM expenses WHERE user_id = %s AND date BETWEEN %s AND %s),
            (SELECT COALESCE(SUM(amount), 0) FROM income WHERE user_id = %s AND date BETWEEN %s AND %s)
    """, (user_id, start_date, end_date, user_id, start_date, end_date))
    sums_row = cursor.fetchone()
    monthly_expenses = float(sums_row[0])
    monthly_income = float(sums_row[1])

    # 2. Category distribution (Expenses only) + per-category transaction count.
    # The count lets us derive the "most used category" below without a 2nd query.
    cursor.execute("""
        SELECT e.category, SUM(e.amount) AS total, COUNT(*) AS cnt, MAX(c.icon) AS category_icon
        FROM expenses e
        LEFT JOIN categories c ON (e.category = c.name AND c.user_id = e.user_id)
        WHERE e.user_id = %s AND e.date BETWEEN %s AND %s
        GROUP BY e.category
        ORDER BY total DESC
    """, (user_id, start_date, end_date))
    category_data = cursor.fetchall()

    # Form lists for Chart.js
    labels = []
    values = []
    for row in category_data:
        labels.append(f"{row['category_icon'] or '📦'} {row['category']}")
        values.append(float(row["total"]))

    labels_json = json.dumps(labels)
    values_json = json.dumps(values)

    # 3. Key Stats: total transactions, average & largest expense (one round trip)
    cursor.execute("""
        SELECT COUNT(*) AS cnt,
               COALESCE(AVG(amount), 0) AS avg_amt,
               COALESCE(MAX(amount), 0) AS max_amt
        FROM expenses WHERE user_id = %s AND date BETWEEN %s AND %s
    """, (user_id, start_date, end_date))
    stats_row = cursor.fetchone()
    total_transactions = stats_row["cnt"]
    avg_expense = float(stats_row["avg_amt"])
    largest_expense = float(stats_row["max_amt"])

    # Highest Spending Day
    cursor.execute("""
        SELECT date, SUM(amount) AS daily_sum 
        FROM expenses 
        WHERE user_id = %s AND date BETWEEN %s AND %s 
        GROUP BY date 
        ORDER BY daily_sum DESC 
        LIMIT 1
    """, (user_id, start_date, end_date))
    highest_day_row = cursor.fetchone()
    if highest_day_row:
        highest_day_val = float(highest_day_row["daily_sum"])
        try:
            d_parsed = datetime.datetime.strptime(highest_day_row["date"], "%Y-%m-%d")
            highest_day_label = d_parsed.strftime("%d %b")
        except ValueError:
            highest_day_label = highest_day_row["date"]
    else:
        highest_day_val = 0
        highest_day_label = "N/A"

    # Average Daily Spend
    try:
        dt_start = datetime.datetime.strptime(start_date, "%Y-%m-%d")
        dt_end = datetime.datetime.strptime(end_date, "%Y-%m-%d")
        days_in_period = max(1, (dt_end - dt_start).days + 1)
    except ValueError:
        days_in_period = 30
    average_daily_spend = float(monthly_expenses) / days_in_period

    # Most Used Category (derived from the distribution query above)
    if category_data:
        most_used_cat_row = max(category_data, key=lambda r: r["cnt"])
        most_used_category = f"{most_used_cat_row['category_icon'] or '📦'} {most_used_cat_row['category']}"
    else:
        most_used_category = "N/A"

    # Most Used Payment Method
    cursor.execute("""
        SELECT e.payment_method, COUNT(*) AS cnt, p.icon AS pm_icon
        FROM expenses e
        LEFT JOIN payment_methods p ON (e.payment_method = p.name AND p.user_id = e.user_id)
        WHERE e.user_id = %s AND e.date BETWEEN %s AND %s
        GROUP BY e.payment_method, p.icon
        ORDER BY cnt DESC
        LIMIT 1
    """, (user_id, start_date, end_date))
    most_used_pm_row = cursor.fetchone()
    if most_used_pm_row:
        most_used_pm = f"{most_used_pm_row['pm_icon'] or '💳'} {most_used_pm_row['payment_method']}"
    else:
        most_used_pm = "N/A"

    # Top Category
    top_category = "N/A"
    if category_data:
        top_category = f"{category_data[0]['category_icon'] or '📦'} {category_data[0]['category']}"

    # Sidebar parameters (single settings round trip)
    currency_symbol, budget_limit = get_currency_symbol_and_budget(user_id)
    budget_percent = min(100, int((monthly_expenses / budget_limit) * 100)) if budget_limit > 0 else 0

    connection.close()

    return render_template(
        "analytics.html",
        monthly_expenses=monthly_expenses,
        monthly_income=monthly_income,
        category_data=category_data,
        labels_json=labels_json,
        values_json=values_json,
        top_category=top_category,
        avg_expense=avg_expense,
        largest_expense=largest_expense,
        current_month_name=range_label,
        filter_range=filter_range,
        selected_month=selected_month,
        total_transactions=total_transactions,
        highest_day_val=highest_day_val,
        highest_day_label=highest_day_label,
        average_daily_spend=average_daily_spend,
        most_used_category=most_used_category,
        most_used_pm=most_used_pm,
        budget_limit=budget_limit,
        budget_percent=budget_percent,
        currency_symbol=currency_symbol
    )


@app.route("/categories")
@login_required
def categories_view():
    user_id = session["user_id"]
    connection = get_db_connection()
    cursor = connection.cursor()

    # Aggregate per-category usage in a single query instead of querying the
    # expenses/income tables once per category.
    cursor.execute("""
        SELECT c.id,
               c.name,
               c.icon,
               COALESCE(e.total_count, 0) + COALESCE(i.total_count, 0) AS count,
               COALESCE(e.total_spent, 0) AS total_spent
        FROM categories c
        LEFT JOIN (
            SELECT category, COUNT(*) AS total_count, SUM(amount) AS total_spent
            FROM expenses
            WHERE user_id = %s
            GROUP BY category
        ) e ON e.category = c.name
        LEFT JOIN (
            SELECT category, COUNT(*) AS total_count
            FROM income
            WHERE user_id = %s
            GROUP BY category
        ) i ON i.category = c.name
        WHERE c.user_id = %s
        ORDER BY c.name ASC
    """, (user_id, user_id, user_id))
    category_rows = cursor.fetchall()

    categories_list = []
    for cat in category_rows:
        categories_list.append({
            "id": cat["id"],
            "name": cat["name"],
            "icon": cat["icon"],
            "count": cat["count"],
            "total_spent": float(cat["total_spent"])
        })

    # Sidebar parameters
    now = datetime.datetime.now()
    current_month_str = now.strftime("%Y-%m")
    cursor.execute("SELECT SUM(amount) AS total FROM expenses WHERE user_id = %s AND substr(date, 1, 7) = %s", (user_id, current_month_str))
    monthly_expenses = cursor.fetchone()[0] or 0
    currency_symbol, budget_limit = get_currency_symbol_and_budget(user_id)
    budget_percent = min(100, int((monthly_expenses / budget_limit) * 100)) if budget_limit > 0 else 0

    connection.close()

    return render_template(
        "categories.html",
        categories=categories_list,
        monthly_expenses=monthly_expenses,
        budget_limit=budget_limit,
        budget_percent=budget_percent,
        currency_symbol=currency_symbol
    )


# =========================================
# REPORT DOWNLOADS AND PDF/CSV EXPORTS
# =========================================

def compute_report_data(user_id, report_type, target_val):
    connection = get_db_connection()
    cursor = connection.cursor()
    
    total_spent = 0
    num_transactions = 0
    expenses = []
    category_breakdown = []
    day_by_day = []
    avg_daily_spend = 0
    highest_spending_day = 0
    raw_cb = []
    raw_expenses = []
    
    if report_type == "daily":
        cursor.execute("""
            SELECT COALESCE(SUM(amount), 0) AS total, COUNT(*) AS cnt
            FROM expenses WHERE user_id = %s AND date = %s
        """, (user_id, target_val))
        daily_agg = cursor.fetchone()
        total_spent = daily_agg["total"]
        num_transactions = daily_agg["cnt"]
        
        cursor.execute("""
            SELECT e.category, SUM(e.amount) AS total, c.icon AS category_icon
            FROM expenses e
            LEFT JOIN categories c ON (e.category = c.name AND c.user_id = e.user_id)
            WHERE e.user_id = %s AND e.date = %s 
            GROUP BY e.category, c.icon 
            ORDER BY total DESC
        """, (user_id, target_val))
        raw_cb = cursor.fetchall()
        
        cursor.execute("""
            SELECT e.date, e.category, e.description, e.payment_method, e.amount, c.icon AS category_icon
            FROM expenses e
            LEFT JOIN categories c ON (e.category = c.name AND c.user_id = e.user_id)
            WHERE e.user_id = %s AND e.date = %s 
            ORDER BY e.id ASC
        """, (user_id, target_val))
        raw_expenses = cursor.fetchall()
        
    elif report_type == "monthly":
        cursor.execute("""
            SELECT COALESCE(SUM(amount), 0) AS total, COUNT(*) AS cnt
            FROM expenses WHERE user_id = %s AND substr(date, 1, 7) = %s
        """, (user_id, target_val))
        monthly_agg = cursor.fetchone()
        total_spent = monthly_agg["total"]
        num_transactions = monthly_agg["cnt"]
        

        
        cursor.execute("""
            SELECT e.category, SUM(e.amount) AS total, c.icon AS category_icon
            FROM expenses e
            LEFT JOIN categories c ON (e.category = c.name AND c.user_id = e.user_id)
            WHERE e.user_id = %s AND substr(e.date, 1, 7) = %s 
            GROUP BY e.category, c.icon 
            ORDER BY total DESC
        """, (user_id, target_val))
        raw_cb = cursor.fetchall()
        
        try:
            year, month = map(int, target_val.split('-'))
            days_in_month = calendar.monthrange(year, month)[1]
        except (ValueError, IndexError):
            days_in_month = 30
            
        avg_daily_spend = float(total_spent) / days_in_month
        
        
        day_sums = {}
        cursor.execute("""
            SELECT date, SUM(amount) AS total 
            FROM expenses 
            WHERE user_id = %s AND substr(date, 1, 7) = %s 
            GROUP BY date
        """, (user_id, target_val))
        for r in cursor.fetchall():
            day_sums[r["date"]] = float(r["total"])
        highest_spending_day = max(day_sums.values(), default=0)
            
        for d in range(1, days_in_month + 1):
            date_str = f"{target_val}-{d:02d}"
            amt = day_sums.get(date_str, 0)
            day_by_day.append({
                "day_label": f"{d} " + calendar.month_abbr[month] if 'month' in locals() else f"{d}",
                "amount": amt
            })
            
        cursor.execute("""
            SELECT e.date, e.category, e.description, e.payment_method, e.amount, c.icon AS category_icon
            FROM expenses e
            LEFT JOIN categories c ON (e.category = c.name AND c.user_id = e.user_id)
            WHERE e.user_id = %s AND substr(e.date, 1, 7) = %s 
            ORDER BY e.date ASC, e.id ASC
        """, (user_id, target_val))
        raw_expenses = cursor.fetchall()
        
    connection.close()

    for row in raw_cb:
        icon = row["category_icon"] or DEFAULT_CATEGORY_ICONS.get(row["category"], "📦")
        category_breakdown.append({
            "category": row["category"],
            "total": float(row["total"]),
            "category_icon": icon
        })

    for exp in raw_expenses:
        icon = exp["category_icon"] or DEFAULT_CATEGORY_ICONS.get(exp["category"], "📦")
        expenses.append({
            "date": exp["date"],
            "category": exp["category"],
            "description": exp["description"],
            "payment_method": exp["payment_method"],
            "amount": float(exp["amount"]),
            "category_icon": icon
        })
    
    return {
        "total_spent": total_spent,
        "num_transactions": num_transactions,
        "expenses": expenses,
        "category_breakdown": category_breakdown,
        "day_by_day": day_by_day,
        "avg_daily_spend": avg_daily_spend,
        "highest_spending_day": highest_spending_day
    }


@app.route("/downloads")
@login_required
def downloads_view():
    user_id = session["user_id"]
    report_type = request.args.get("type", "monthly").strip().lower()
    target_date = request.args.get("date", "").strip()
    target_month = request.args.get("month", "").strip()
    
    now = datetime.datetime.now()
    if not target_date:
        target_date = now.strftime("%Y-%m-%d")
    if not target_month:
        target_month = now.strftime("%Y-%m")
        
    target_val = target_date if report_type == "daily" else target_month
    report_data = compute_report_data(user_id, report_type, target_val)
    
    has_data = report_data["num_transactions"] > 0
    
    if report_type == "daily":
        try:
            d_parsed = datetime.datetime.strptime(target_date, "%Y-%m-%d")
            report_title = d_parsed.strftime("%d %b %Y")
        except ValueError:
            report_title = target_date
    else:
        try:
            d_parsed = datetime.datetime.strptime(target_month, "%Y-%m")
            report_title = d_parsed.strftime("%B %Y")
        except ValueError:
            report_title = target_month

    connection = get_db_connection()
    cursor = connection.cursor()
    current_month_str = now.strftime("%Y-%m")
    cursor.execute("SELECT SUM(amount) AS total FROM expenses WHERE user_id = %s AND substr(date, 1, 7) = %s", (user_id, current_month_str))
    monthly_expenses = cursor.fetchone()[0] or 0
    connection.close()

    currency_symbol, budget_limit = get_currency_symbol_and_budget(user_id)
    budget_percent = min(100, int((monthly_expenses / budget_limit) * 100)) if budget_limit > 0 else 0

    return render_template(
        "downloads.html",
        report_type=report_type,
        target_date=target_date,
        target_month=target_month,
        report_data=report_data,
        has_data=has_data,
        report_title=report_title,
        monthly_expenses=monthly_expenses,
        budget_limit=budget_limit,
        budget_percent=budget_percent,
        currency_symbol=currency_symbol
    )


@app.route("/downloads/csv")
@login_required
def downloads_csv():
    user_id = session["user_id"]
    report_type = request.args.get("type", "monthly").strip().lower()
    target_date = request.args.get("date", "").strip()
    target_month = request.args.get("month", "").strip()
    
    now = datetime.datetime.now()
    if not target_date:
        target_date = now.strftime("%Y-%m-%d")
    if not target_month:
        target_month = now.strftime("%Y-%m")
        
    target_val = target_date if report_type == "daily" else target_month
    report_data = compute_report_data(user_id, report_type, target_val)
    
    if report_data["num_transactions"] == 0:
        return "No data available for this period.", 400
        
    si = io.StringIO()
    cw = csv.writer(si)
    cw.writerow(["Date", "Category", "Description", "Payment Method", "Amount"])
    for expense in report_data["expenses"]:
        cw.writerow([
            expense["date"],
            expense["category"],
            expense["description"] or "",
            expense["payment_method"],
            f"{float(expense['amount']):.2f}"
        ])
        
    output = make_response(si.getvalue())
    output.headers["Content-Disposition"] = f"attachment; filename=report_{report_type}_{target_val}.csv"
    output.headers["Content-type"] = "text/csv"
    return output


@app.route("/downloads/xlsx")
@login_required
def downloads_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    user_id = session["user_id"]
    report_type = request.args.get("type", "monthly").strip().lower()
    target_date = request.args.get("date", "").strip()
    target_month = request.args.get("month", "").strip()
    
    now = datetime.datetime.now()
    if not target_date:
        target_date = now.strftime("%Y-%m-%d")
    if not target_month:
        target_month = now.strftime("%Y-%m")
        
    target_val = target_date if report_type == "daily" else target_month
    report_data = compute_report_data(user_id, report_type, target_val)
    
    if report_data["num_transactions"] == 0:
        return "No data available for this period.", 400

    wb = Workbook()
    ws = wb.active
    ws.title = "Report"

    headers = ["Date", "Category", "Description", "Payment Method", "Amount"]
    widths = [15, 22, 45, 26, 18]

    header_fill = PatternFill("solid", fgColor="0C110F")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    band_fill = PatternFill("solid", fgColor="F2F8F5")
    thin = Side(style="thin", color="D7E3DC")
    cell_border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col_idx, (header, width) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = cell_border
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 22

    for row_idx, expense in enumerate(report_data["expenses"], start=2):
        values = [
            expense["date"],
            expense["category"],
            expense["description"] or "",
            expense["payment_method"],
            float(expense["amount"]),
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = cell_border
            if col_idx == 5:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "#,##0.00"
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            if row_idx % 2 == 0:
                cell.fill = band_fill

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:E{len(report_data['expenses']) + 1}"

    buffer = io.BytesIO()
    wb.save(buffer)
    output = make_response(buffer.getvalue())
    output.headers["Content-Disposition"] = f"attachment; filename=report_{report_type}_{target_val}.xlsx"
    output.headers["Content-Type"] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return output


def _get_pdf_class():
    from fpdf import FPDF
    class ExpenseReportPDF(FPDF):
        PAGE_W = 210.0
        MARGIN = 12.0
        CONTENT_TOP = 40.0
        CONTENT_W = PAGE_W - 2 * MARGIN

        BG = (12, 17, 15)
        CARD = (18, 24, 21)
        ROW_ALT = (17, 23, 20)
        GREEN = (32, 200, 120)
        WHITE = (238, 245, 241)
        MUTED = (150, 163, 156)
        DIM = (105, 120, 112)
        BORDER = (38, 52, 45)

        def __init__(self, currency_symbol="Rs", *args, **kwargs):
            super().__init__(*args, **kwargs)
            self.currency_symbol = currency_symbol

        def normalize_text(self, text):
            if isinstance(text, bytes):
                text = text.decode("latin-1")
            try:
                text.encode("latin-1")
            except UnicodeEncodeError:
                text = text.replace("₹", "Rs ").replace("€", "EUR ")
            return super().normalize_text(text)

        def header(self):
            self.set_fill_color(*self.BG)
            self.rect(0, 0, self.PAGE_W, self.h, "F")
            self.set_fill_color(*self.GREEN)
            self.rect(0, 0, self.PAGE_W, 2.2, "F")
            if self.page_no() == 1:
                self.set_fill_color(*self.GREEN)
                self.rect(self.MARGIN, 9.5, 3, 13, "F")
                self.set_text_color(*self.GREEN)
                self.set_font("helvetica", "B", 20)
                self.set_xy(self.MARGIN + 7, 8.5)
                self.cell(120, 8, "EXPENSE MANAGER", align="L")
                self.set_text_color(*self.MUTED)
                self.set_font("helvetica", "I", 9)
                self.set_xy(self.MARGIN + 7, 17.5)
                self.cell(120, 5, "Expense Reports Log Summary", align="L")
                self.set_fill_color(*self.GREEN)
                self.rect(self.PAGE_W - self.MARGIN - 38, 10.5, 38, 10, "F")
                self.set_text_color(*self.BG)
                self.set_font("helvetica", "B", 9)
                self.set_xy(self.PAGE_W - self.MARGIN - 38, 12.5)
                self.cell(38, 6, "REPORT", align="C")
                self.set_draw_color(*self.GREEN)
                self.set_line_width(0.6)
                self.line(self.MARGIN, 34, self.PAGE_W - self.MARGIN, 34)
                self.set_y(self.CONTENT_TOP)
            else:
                self.set_font("helvetica", "I", 8)
                self.set_text_color(*self.MUTED)
                self.set_xy(self.MARGIN, 4)
                self.cell(self.CONTENT_W, 5, "EXPENSE MANAGER - Continued", align="L")
                self.set_draw_color(*self.GREEN)
                self.set_line_width(0.5)
                self.line(self.MARGIN, 12.5, self.PAGE_W - self.MARGIN, 12.5)
                self.set_y(14)

        def footer(self):
            self.set_y(-14)
            self.set_fill_color(*self.BG)
            self.rect(0, self.h - 14, self.PAGE_W, 14, "F")
            self.set_draw_color(*self.GREEN)
            self.set_line_width(0.3)
            self.line(self.MARGIN, self.h - 14, self.PAGE_W - self.MARGIN, self.h - 14)
            self.set_text_color(*self.MUTED)
            self.set_font("helvetica", "I", 8)
            self.cell(0, 10, f"Page {self.page_no()} | Generated by Expense Manager", align="C")

        def _ensure_space(self, height):
            if self.get_y() + height > self.page_break_trigger - 3:
                self.add_page()
                return True
            return False

        def section_heading(self, text, continued=False):
            self.ln(2)
            y0 = self.get_y()
            self.set_fill_color(*self.GREEN)
            self.rect(self.MARGIN, y0 + 1.2, 3.2, 6, "F")
            self.set_x(self.MARGIN + 7)
            self.set_font("helvetica", "B", 12)
            self.set_text_color(*self.GREEN)
            label = text + (" (continued)" if continued else "")
            self.cell(self.CONTENT_W - 7, 8, label, ln=True)
            self.set_draw_color(*self.GREEN)
            self.set_line_width(0.5)
            self.line(self.MARGIN, self.get_y() + 0.6, self.PAGE_W - self.MARGIN, self.get_y() + 0.6)
            self.ln(5)

        def stat_card(self, x, y, w, h, label, value):
            self.set_fill_color(*self.CARD)
            self.set_draw_color(*self.GREEN)
            self.set_line_width(0.3)
            self.rect(x, y, w, h, "DF")
            self.set_fill_color(*self.GREEN)
            self.rect(x, y, 1.4, h, "F")
            self.set_font("helvetica", "B", 7)
            self.set_text_color(*self.GREEN)
            self.set_xy(x + 4.5, y + 3.5)
            self.cell(w - 8, 3.5, label, align="L")
            self.set_font("helvetica", "B", 14)
            self.set_text_color(*self.WHITE)
            self.set_xy(x + 4.5, y + 8)
            self.cell(w - 8, 6, value, align="L")

        def category_row(self, index, category, amount):
            y0 = self.get_y()
            self.set_fill_color(*(self.ROW_ALT if index % 2 else self.CARD))
            self.set_draw_color(*self.BORDER)
            self.set_line_width(0.2)
            self.rect(self.MARGIN, y0, self.CONTENT_W, 8, "DF")
            self.set_fill_color(*self.GREEN)
            self.rect(self.MARGIN, y0, 1.2, 8, "F")
            self.set_font("helvetica", "", 9.5)
            self.set_text_color(*self.WHITE)
            self.set_xy(self.MARGIN + 4, y0 + 1.5)
            self.cell(self.CONTENT_W - 54, 5, self.fit(category, self.CONTENT_W - 54), align="L")
            self.set_font("helvetica", "B", 9.5)
            self.set_text_color(*self.GREEN)
            self.set_xy(self.MARGIN + 50, y0 + 1.5)
            self.cell(self.CONTENT_W - 54, 5, amount, align="R")
            self.set_y(y0 + 8)

        def day_cell(self, x, y, w, h, day_label, amount):
            self.set_fill_color(*self.CARD)
            self.set_draw_color(*self.GREEN)
            self.set_line_width(0.25)
            self.rect(x, y, w, h, "DF")
            self.set_font("helvetica", "B", 7)
            self.set_text_color(*self.GREEN)
            self.set_xy(x + 3, y + 2)
            self.cell(w - 6, 4, day_label, align="C")
            self.set_font("helvetica", "B", 9)
            self.set_text_color(*self.WHITE)
            self.set_xy(x + 3, y + 6.5)
            self.cell(w - 6, 4.5, amount, align="C")

        def table_header_row(self):
            y0 = self.get_y()
            self.set_fill_color(*self.GREEN)
            self.set_draw_color(*self.GREEN)
            self.set_line_width(0.3)
            self.rect(self.MARGIN, y0, self.CONTENT_W, 8, "DF")
            self.set_font("helvetica", "B", 8)
            self.set_text_color(*self.BG)
            cols = [("Date", 24, "L"), ("Category", 32, "L"), ("Description", 68, "L"), ("Payment Method", 32, "L"), ("Amount", 30, "R")]
            x = self.MARGIN
            for txt, w, align in cols:
                self.set_xy(x + (3 if align == "L" else 0), y0 + 2)
                self.cell(w, 4, txt, align=align)
                x += w
            self.set_y(y0 + 8)

        def expense_row(self, index, expense, currency_symbol):
            y0 = self.get_y()
            self.set_fill_color(*(self.ROW_ALT if index % 2 else self.BG))
            self.set_draw_color(*self.BORDER)
            self.set_line_width(0.2)
            self.rect(self.MARGIN, y0, self.CONTENT_W, 7.5, "DF")
            self.set_font("helvetica", "", 8.5)
            self.set_text_color(*self.WHITE)
            self.set_xy(self.MARGIN + 3, y0 + 1.7)
            self.cell(24, 4, expense["date"], align="L")
            self.set_xy(self.MARGIN + 24 + 3, y0 + 1.7)
            self.cell(32 - 1, 4, self.fit(expense["category"], 28), align="L")
            self.set_text_color(*self.MUTED)
            self.set_xy(self.MARGIN + 24 + 32 + 3, y0 + 1.7)
            self.cell(68 - 1, 4, self.fit(expense["description"] or "", 66), align="L")
            self.set_text_color(*self.WHITE)
            self.set_xy(self.MARGIN + 24 + 32 + 68 + 3, y0 + 1.7)
            self.cell(32 - 1, 4, self.fit(expense["payment_method"], 30), align="L")
            self.set_font("helvetica", "B", 8.5)
            self.set_text_color(*self.GREEN)
            self.set_xy(self.MARGIN + 24 + 32 + 68 + 32, y0 + 1.7)
            self.cell(30 - 2, 4, f"{currency_symbol}{float(expense['amount']):.2f}", align="R")
            self.set_y(y0 + 7.5)

        def fit(self, text, width):
            text = text or ""
            if self.get_string_width(text) <= width:
                return text
            result = text
            while result and self.get_string_width(result + "...") > width:
                result = result[:-1]
            return result + "..."
    return ExpenseReportPDF


@app.route("/downloads/pdf")
@login_required
def downloads_pdf():
    user_id = session["user_id"]
    report_type = request.args.get("type", "monthly").strip().lower()
    target_date = request.args.get("date", "").strip()
    target_month = request.args.get("month", "").strip()
    
    now = datetime.datetime.now()
    if not target_date:
        target_date = now.strftime("%Y-%m-%d")
    if not target_month:
        target_month = now.strftime("%Y-%m")
        
    target_val = target_date if report_type == "daily" else target_month
    report_data = compute_report_data(user_id, report_type, target_val)
    
    if report_data["num_transactions"] == 0:
        return "No data available for this period.", 400
        
    _, currency_symbol = get_currency_info(user_id)
    
    if report_type == "daily":
        try:
            d_parsed = datetime.datetime.strptime(target_date, "%Y-%m-%d")
            report_title = d_parsed.strftime("%d %b %Y")
        except ValueError:
            report_title = target_date
        report_label = f"Daily Expense Report - {report_title}"
    else:
        try:
            d_parsed = datetime.datetime.strptime(target_month, "%Y-%m")
            report_title = d_parsed.strftime("%B %Y")
        except ValueError:
            report_title = target_month
        report_label = f"Monthly Expense Report - {report_title}"

    pdf = _get_pdf_class()(currency_symbol=currency_symbol)
    pdf.set_margins(pdf.MARGIN, pdf.CONTENT_TOP, pdf.MARGIN)
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=14)

    pdf._ensure_space(20)
    y0 = pdf.get_y()
    pdf.set_fill_color(*pdf.GREEN)
    pdf.rect(pdf.MARGIN, y0 + 1.2, 3.2, 6, "F")
    pdf.set_x(pdf.MARGIN + 7)
    pdf.set_font("helvetica", "B", 15)
    pdf.set_text_color(*pdf.GREEN)
    pdf.cell(pdf.CONTENT_W - 7, 8, report_label, ln=True)
    pdf.set_x(pdf.MARGIN + 7)
    pdf.set_font("helvetica", "", 9)
    pdf.set_text_color(*pdf.MUTED)
    pdf.cell(pdf.CONTENT_W - 7, 5, f"Report Date: {now.strftime('%d %b %Y %H:%M')}  |  Currency: {currency_symbol}", ln=True)
    pdf.ln(6)

    card_h = 24
    gap = 4
    if report_type == "daily":
        pdf._ensure_space(card_h)
        card_w = (pdf.CONTENT_W - gap) / 2
        x0 = pdf.MARGIN
        y = pdf.get_y()
        pdf.stat_card(x0, y, card_w, card_h, "Total Spent",
                      f"{currency_symbol}{float(report_data['total_spent']):.2f}")
        pdf.stat_card(x0 + card_w + gap, y, card_w, card_h, "Transactions",
                      f"  {report_data['num_transactions']}")
        pdf.set_y(y + card_h)
    else:
        pdf._ensure_space(card_h)
        card_w = (pdf.CONTENT_W - 3 * gap) / 4
        x0 = pdf.MARGIN
        y = pdf.get_y()
        cards = [
            ("Spent", f"{currency_symbol}{float(report_data['total_spent']):.0f}"),
            ("Transactions", f"  {report_data['num_transactions']}"),
            ("Avg Daily", f"{currency_symbol}{float(report_data['avg_daily_spend']):.2f}"),
            ("Max Day", f"{currency_symbol}{float(report_data['highest_spending_day']):.0f}"),
        ]
        for i, (label, value) in enumerate(cards):
            pdf.stat_card(x0 + i * (card_w + gap), y, card_w, card_h, label, value)
        pdf.set_y(y + card_h)
    pdf.ln(3)

    pdf.section_heading("Category Breakdown Summary")
    for i, row in enumerate(report_data["category_breakdown"]):
        if pdf._ensure_space(8):
            pdf.section_heading("Category Breakdown Summary", continued=True)
        pdf.category_row(i, row["category"], f"{currency_symbol}{float(row['total']):.2f}")
    pdf.ln(4)

    if report_type == "monthly":
        pdf.section_heading("Day-by-Day Spending Summary")
        items = list(report_data["day_by_day"])
        col_width = (pdf.CONTENT_W - 3 * gap) / 4
        cell_h = 13
        row_h = cell_h + 1
        first_row = True
        y_row = pdf.get_y()
        for idx in range(0, len(items), 4):
            if y_row + row_h > pdf.page_break_trigger - 3:
                pdf.add_page()
                pdf.section_heading("Day-by-Day Spending Summary", continued=not first_row)
                y_row = pdf.get_y()
            chunk = items[idx:idx + 4]
            x = pdf.MARGIN
            for item in chunk:
                pdf.day_cell(x, y_row, col_width, cell_h,
                             item["day_label"],
                             f"{currency_symbol}{float(item['amount']):.0f}")
                x += col_width + gap
            y_row += row_h
            first_row = False
        pdf.set_y(y_row)
        pdf.ln(2)

    row_h = 7.5
    if report_data["expenses"] and pdf.get_y() + 15 + 8 + row_h > pdf.page_break_trigger - 3:
        pdf.add_page()
    pdf.section_heading("Detailed Expense Log List")
    pdf.table_header_row()
    for i, exp in enumerate(report_data["expenses"]):
        if pdf._ensure_space(row_h):
            pdf.section_heading("Detailed Expense Log List", continued=True)
            pdf.table_header_row()
        pdf.expense_row(i, exp, currency_symbol)
    pdf.ln(2)

    pdf_bytes = bytes(pdf.output())
    response = make_response(pdf_bytes)
    response.headers["Content-Disposition"] = f"attachment; filename=report_{report_type}_{target_val}.pdf"
    response.headers["Content-Type"] = "application/pdf"
    return response


if __name__ == "__main__":
    app.run(debug=True)